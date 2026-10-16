#!/usr/bin/env python3
"""Ubah CC_Expense_Tracker.xlsx (Google Sheets) jadi JSON untuk diimpor app.

Pemakaian:
  python3 tool/xlsx_to_json.py CC_Expense_Tracker.xlsx import_data.json \
      [--cutoff 4] [--cicilan-paid "Kursi=2,Vacuum cleaner=1"] \
      [--subs "Claude=349000@3,Hevy=45000@25"] [--sub-status "Yellow=dibatalkan"] \
      [--fix-total "O19=164325"] [--drop-split "O20=Daniel:2"]

Format JSON (version 1) dibaca oleh SheetJsonImporter di
lib/data/import/. Nama enum memakai nama enum Dart.

Catatan konversi:
  * Kolom "Bulan" di sheet Transaksi -> bulanTagihan (bisa beda dari Tanggal).
  * Transaksi tanpa Tanggal memakai tanggal 1 pada Bulan-nya.
  * "Cicilan ke-" di sheet = cicilan bulan ini (termasuk yang belum dibayar);
    di app cicilanKe = jumlah yang SUDAH dibayar. Default: sama dengan sheet
    kalau "Status Bulan Ini" berisi Sudah Dibayar/Paid, selain itu dikurangi 1.
    Override dengan --cicilan-paid.
  * Sheet Subscription tidak punya harga/tanggal tagihan, jadi tanpa --subs
    subscription diimpor sebagai harga variable (tanpa auto-generate).
    --subs "Nama=harga@tanggal" mengisi harga tetap + hari tagihan (auto-generate
    aktif untuk yang Aktif); --sub-status mengganti status (aktif|nonaktif|dibatalkan).
  * Transaksi Subscription ditautkan ke layanan: (1) Deskripsi memuat nama
    layanan; (2) nominal sama dan hari <= 3 dari tanggal tagihan; (3) hari saja,
    kalau kandidatnya tunggal. Sisanya dibiarkan tidak tertaut.
  * Tanggal tagihan berikutnya = bulan setelah transaksi tertaut terakhir.
  * Koreksi data (tercatat di output sebagai [koreksi]):
    --fix-total "ID Order=total" mengganti Total Tagihan Group Order;
    --drop-split "ID Order=Nama:ke-N" membuang kemunculan ke-N nama itu
    (ke-1 kalau ":N" tidak ditulis), untuk baris peserta yang salah input.
  * Group Order yang jumlah split-nya beda dari totalnya diberi peringatan.
"""
import argparse
import datetime
import json
import sys

import openpyxl

JENIS = {'Personal': 'personal', 'Group Order': 'groupOrder', 'Subscription': 'subscription',
         'Pembayaran CC': 'pembayaranCC', 'Lainnya': 'lainnya'}
SUMBER = {'CC': 'cc', 'Saldo': 'saldo'}
STATUS = {'Paid': 'paid', 'Paid Temp': 'paidTemp', 'Belum Dibayar': 'belumDibayar'}
SIKLUS = {'Bulanan': 'bulanan', 'Tahunan': 'tahunan'}
STATUS_SUB = {'Aktif': 'aktif', 'Nonaktif': 'nonaktif', 'Dibatalkan': 'dibatalkan'}
PAID_WORDS = {'sudah dibayar', 'paid', 'lunas', 'sudah bayar'}

warnings = []


def warn(msg):
    warnings.append(msg)


def to_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    return None


def month_str(v):
    """Kolom Bulan: bisa datetime atau teks 'YYYY-MM'."""
    d = to_date(v)
    if d:
        return d.strftime('%Y-%m')
    return str(v).strip() if v else None


def add_months(d, n):
    total = d.year * 12 + d.month - 1 + n
    y, m = divmod(total, 12)
    return datetime.date(y, m + 1, 1)


def as_int(v):
    return None if v is None or v == '' else int(round(float(v)))


def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def mapped(table, value, what, row):
    key = clean(value)
    if key not in table:
        sys.exit(f'{what} tidak dikenal di {row}: {value!r}')
    return table[key]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx')
    ap.add_argument('out')
    ap.add_argument('--cutoff', type=int, default=4, help='tanggal statement CC tutup buku (0 = bulan kalender)')
    ap.add_argument('--cicilan-paid', default='', help='override jumlah cicilan yang sudah dibayar: "Nama=2,Lain=1"')
    ap.add_argument('--subs', default='', help='"Nama=harga@tanggal,..." harga tetap + hari tagihan')
    ap.add_argument('--fix-total', default='', help='"O19=164325,..." koreksi Total Tagihan Group Order')
    ap.add_argument('--drop-split', default='', help='"O20=Daniel:2,..." buang baris peserta yang salah input')
    ap.add_argument('--sub-status', default='', help='"Nama=dibatalkan,..."')
    ap.add_argument('--today', default=datetime.date.today().isoformat())
    args = ap.parse_args()
    today = datetime.date.fromisoformat(args.today)
    overrides = {}
    for part in filter(None, (p.strip() for p in args.cicilan_paid.split(','))):
        name, _, n = part.rpartition('=')
        overrides[name.strip().lower()] = int(n)

    sub_cfg = {}
    for part in filter(None, (p.strip() for p in args.subs.split(','))):
        name, _, rest = part.rpartition('=')
        price, _, day = rest.partition('@')
        sub_cfg[name.strip().lower()] = (int(price), int(day))
    status_cfg = {}
    for part in filter(None, (p.strip() for p in args.sub_status.split(','))):
        name, _, st = part.rpartition('=')
        if st.strip() not in STATUS_SUB.values():
            sys.exit(f'--sub-status: status tidak dikenal {st!r}')
        status_cfg[name.strip().lower()] = st.strip()

    fix_total = {}
    for part in filter(None, (p.strip() for p in args.fix_total.split(','))):
        order, _, value = part.partition('=')
        fix_total[order.strip()] = int(value)
    drop_split = {}
    for part in filter(None, (p.strip() for p in args.drop_split.split(','))):
        order, _, rest = part.partition('=')
        name, _, nth = rest.partition(':')
        drop_split.setdefault(order.strip(), []).append((name.strip(), int(nth or 1)))
    fixes = []

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)

    # --- split per ID Order
    splits_by_order = {}
    for r in wb['Split Group Order'].iter_rows(min_row=2, values_only=True):
        if r[0]:
            splits_by_order.setdefault(str(r[0]).strip(), []).append(r)

    # --- subscription registry
    subs = []
    for r in wb['Subscription'].iter_rows(min_row=2, values_only=True):
        if not r[0] or r[1] not in SIKLUS:  # lewati baris total/catatan
            continue
        subs.append({'namaLayanan': clean(r[0]), 'hargaTetap': None,
                     'siklusBilling': SIKLUS[r[1]],
                     'status': mapped(STATUS_SUB, r[3], 'Status subscription', r[0]),
                     'catatan': clean(r[4]), '_last': None})
    known = {s['namaLayanan'].lower() for s in subs}
    for cfg_name in list(sub_cfg) + list(status_cfg):
        if cfg_name not in known:
            sys.exit(f'--subs/--sub-status: layanan {cfg_name!r} tidak ada di sheet Subscription')
    for s in subs:
        key = s['namaLayanan'].lower()
        if key in sub_cfg:
            s['hargaTetap'], s['hariTagihan'] = sub_cfg[key]
        if key in status_cfg:
            s['status'] = status_cfg[key]

    # --- transaksi
    transaksis = []
    unlinked_subs = []
    pending_links = []
    for r in wb['Transaksi'].iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        no = int(r[0])
        row = f'Transaksi No {no}'
        bulan = month_str(r[1])
        if not bulan:
            sys.exit(f'{row}: kolom Bulan kosong')
        tanggal = to_date(r[2]) or datetime.date(int(bulan[:4]), int(bulan[5:7]), 1)
        jenis = mapped(JENIS, r[4], 'Jenis', row)
        deskripsi = clean(r[3]) or '(tanpa deskripsi)'
        total = as_int(r[6])
        t = {
            'tanggal': tanggal.isoformat(), 'bulanTagihan': bulan, 'deskripsi': deskripsi,
            'jenis': jenis, 'sumberBayar': mapped(SUMBER, r[5], 'Sumber Bayar', row),
            'totalTagihan': total, 'merchant': None, 'ongkirFee': 0,
            'status': mapped(STATUS, r[10], 'Status', row), 'catatan': clean(r[11]),
            'subscription': None, 'sumberEkstraksi': 'manual', 'splits': [],
        }
        if jenis == 'groupOrder':
            order_id = str(r[7]).strip() if r[7] else None
            rows = list(splits_by_order.get(order_id, [])) if order_id else []
            for drop_name, nth in drop_split.get(order_id, []):
                hits = [x for x in rows if clean(x[4]) == drop_name]
                if len(hits) < nth:
                    sys.exit(f'--drop-split {order_id}: tidak ada {drop_name!r} ke-{nth}')
                rows.remove(hits[nth - 1])
                fixes.append(f'{order_id}: baris peserta {drop_name!r} ke-{nth} dibuang')
            if order_id in fix_total:
                fixes.append(f'{order_id}: Total Tagihan {total} -> {fix_total[order_id]}')
                total = fix_total[order_id]
                t['totalTagihan'] = total
            if not rows:
                warn(f'{row} ({deskripsi}): Group Order tanpa baris split')
            else:
                n = len(rows)
                t['merchant'] = clean(rows[0][3])
                t['ongkirFee'] = as_int(rows[0][6]) or 0
                if (as_int(rows[0][5]) != total):
                    warn(f'{row}: Total Order split ({rows[0][5]}) != Total Tagihan ({total})')
                names = [clean(x[4]) for x in rows]
                if len(set(names)) != n:
                    warn(f'{row}: nama peserta kembar {sorted(set(x for x in names if names.count(x) > 1))}')
                for x in rows:
                    manual = as_int(x[10])
                    sheet_amount = float(x[10]) if x[10] not in (None, '') else (
                        (float(x[5] or 0) - float(x[7] or 0) + float(x[6] or 0)) / n)
                    hm = manual
                    if hm is None and abs(sheet_amount - total // n) >= 1:
                        hm = int(round(sheet_amount))  # rumus sheet beda dari total/N (ada ongkir/diskon)
                        warn(f'{row}: {clean(x[4])} diberi hargaManual {hm} agar sama dengan sheet')
                    t['splits'].append({'nama': names[rows.index(x)], 'hargaManual': hm,
                                        'status': mapped(STATUS, x[12], 'Status split', row),
                                        'catatan': clean(x[13])})
        if jenis == 'subscription':
            match = next((s for s in subs if s['namaLayanan'].lower() in deskripsi.lower()), None)
            if match:
                t['subscription'] = match['namaLayanan']
                if not match['_last'] or tanggal > match['_last']:
                    match['_last'] = tanggal
            else:
                pending_links.append((t, tanggal, total, deskripsi))
        transaksis.append(t)

    # Tahap 2: tautkan sisanya lewat nominal + hari, lalu hari saja.
    link_report = []
    configured = [s for s in subs if s.get('hariTagihan')]
    for t, tanggal, total, deskripsi in pending_links:
        near = [s for s in configured if abs(tanggal.day - s['hariTagihan']) <= 3]
        exact = [s for s in near if s['hargaTetap'] == total]
        match, how = (exact[0], 'nominal+hari') if len(exact) == 1 else (
            (near[0], 'hari saja') if not exact and len(near) == 1 else (None, None))
        if match:
            t['subscription'] = match['namaLayanan']
            if not match['_last'] or tanggal > match['_last']:
                match['_last'] = tanggal
            link_report.append(f'{tanggal} {deskripsi} Rp{total} -> {match["namaLayanan"]} ({how})')
        else:
            unlinked_subs.append(f'{tanggal} {deskripsi} Rp{total}')

    for t in transaksis:
        if t['jenis'] != 'groupOrder' or not t['splits'] or all(x['hargaManual'] is None for x in t['splits']):
            continue
        n = len(t['splits'])
        amounts = [x['hargaManual'] if x['hargaManual'] is not None else t['totalTagihan'] // n for x in t['splits']]
        if sum(amounts) != t['totalTagihan']:
            warn(f"{t['tanggal']} {t['deskripsi']}: jumlah split {sum(amounts)} != total {t['totalTagihan']} "
                 f"(selisih {sum(amounts) - t['totalTagihan']:+d})")

    for s in subs:
        base = s.pop('_last')
        if base:
            first = add_months(base, 1)
            anchor = s.get('hariTagihan') or base.day
            last_day = (add_months(first, 1) - datetime.timedelta(days=1)).day
            nxt = first.replace(day=min(anchor, last_day))
        else:
            nxt = add_months(today, 1)
        s['tanggalTagihanBerikutnya'] = nxt.isoformat()

    # --- cicilan
    cicilans = []
    cicilan_report = []
    for r in wb['Cicilan'].iter_rows(min_row=2, values_only=True):
        if not isinstance(r[1], (int, float)) or not r[0]:
            continue  # baris catatan
        nama, tenor, ke = clean(r[0]), as_int(r[2]), as_int(r[3]) or 0
        mulai = to_date(r[7])
        if not mulai:
            sys.exit(f'Cicilan {nama}: Tanggal Mulai kosong')
        status_bulan = (clean(r[9]) or '').lower()
        default_paid = ke if status_bulan in PAID_WORDS else max(ke - 1, 0)
        paid = overrides.get(nama.lower(), default_paid)
        paid = max(0, min(paid, tenor))
        cicilan_report.append(f'{nama}: sheet "Cicilan ke-" {ke}, status bulan ini "{r[9]}" -> cicilanKe (sudah dibayar) {paid}'
                              + (' [override]' if nama.lower() in overrides else ''))
        cicilans.append({
            'namaCicilan': nama, 'totalHargaBarang': as_int(r[1]), 'tenor': tenor, 'cicilanKe': paid,
            'jumlahPerBulan': as_int(r[4]), 'tanggalMulai': mulai.isoformat(),
            'tanggalBayarTerakhir': add_months(mulai, paid - 1).isoformat() if paid > 0 else None,
            'catatan': clean(r[10]),
        })

    d = wb['Dashboard']
    out = {
        'version': 1,
        'settings': {'limitCC': as_int(d['B6'].value), 'saldoOutstandingSebelumTracker': as_int(d['B7'].value),
                     'tanggalCutoffTagihan': args.cutoff},
        'subscriptions': subs, 'cicilans': cicilans, 'transaksis': transaksis,
    }
    with open(args.out, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=1)

    n_splits = sum(len(t['splits']) for t in transaksis)
    print(f'Ditulis {args.out}: {len(transaksis)} transaksi, {n_splits} split, '
          f'{len(cicilans)} cicilan, {len(subs)} subscription')
    print('\nCicilan:')
    for line in cicilan_report:
        print('  ' + line)
    print('\nSubscription (harga tetap / hari / status / tagihan berikutnya):')
    for s in subs:
        print(f"  {s['namaLayanan']}: {s['hargaTetap']} / {s.get('hariTagihan')} / {s['status']} / {s['tanggalTagihanBerikutnya']}")
    for line in link_report:
        print('  ditautkan: ' + line)
    linked = sum(1 for t in transaksis if t['subscription'])
    print(f'\nSubscription: {linked} transaksi tertaut ke layanan; {len(unlinked_subs)} belum tertaut:')
    for u in unlinked_subs:
        print('  ' + u)
    if fixes:
        print('\nKoreksi:')
        for f in fixes:
            print('  [koreksi] ' + f)
    unused = [o for o in list(fix_total) + list(drop_split) if not any(f.startswith(o + ':') for f in fixes)]
    if unused:
        sys.exit(f'Koreksi tidak terpakai (ID Order tidak ditemukan?): {unused}')
    if warnings:
        print('\nPeringatan:')
        for w in warnings:
            print('  - ' + w)


if __name__ == '__main__':
    main()
